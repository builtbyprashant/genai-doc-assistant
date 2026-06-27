"""Tests for app/services/document_loader.py and storage.py (Task 3).

Covers DECISIONS D-2 (validate raw bytes, then parse per-format) and C-C
(SHA-256 dedup). Binary fixtures (PDF/docx/xlsx) are built programmatically so
the suite needs no checked-in sample files.
"""

from __future__ import annotations

import hashlib
import io

import pandas as pd
import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.core import config
from app.services import storage
from app.services.document_loader import (
    DocumentError,
    content_hash,
    load_document,
)


# ── fixtures / builders ──────────────────────────────────────────────────────

def _text_pdf(text: str = "Hello PDF World") -> bytes:
    """A minimal valid PDF with a real text layer (correct xref offsets)."""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = b"BT /F1 24 Tf 72 720 Td (" + text.encode("latin-1") + b") Tj ET"
    objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\n"
    out += b"startxref\n" + str(xref_pos).encode() + b"\n%%EOF"
    return bytes(out)


def _blank_pdf(password: str | None = None) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if password:
        writer.encrypt(password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("ICU transfer policy details for discharge.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    frame = pd.DataFrame({"name": ["John", "Mary"], "age": [45, 62]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Patients")
    return buf.getvalue()


def _set_upload_limit_mb(monkeypatch, value: str) -> None:
    monkeypatch.setenv("UPLOAD_SIZE_LIMIT_MB", value)
    config.get_settings.cache_clear()


# ── happy path: all 8 formats parse ───────────────────────────────────────────

def test_loads_txt():
    out = load_document("a.txt", b"Discharge requires a completed checklist.")
    assert "Discharge" in out["text"]


def test_loads_md():
    out = load_document("a.md", b"# Heading\n\nSome **markdown** body text here.")
    assert "markdown" in out["text"]


def test_loads_csv_with_header():
    out = load_document("a.csv", b"name,age\nJohn,45\nMary,62\nBob,30\n")
    assert "name: John" in out["text"]
    assert out["warnings"] == []


def test_loads_csv_without_header_warns():
    out = load_document("a.csv", b"1,2,3\n4,5,6\n7,8,9\n")
    assert "Col_1" in out["text"]
    assert any("headers" in w.lower() for w in out["warnings"])


def test_loads_json_flattened():
    out = load_document("a.json", b'{"patient": {"name": "John", "age": 45}}')
    assert "patient.name: John" in out["text"]


def test_loads_yaml_flattened():
    out = load_document("a.yaml", b"patient:\n  name: John\n  age: 45\n")
    assert "patient.name: John" in out["text"]


def test_loads_xlsx_with_sheet_prefix():
    out = load_document("a.xlsx", _xlsx_bytes())
    assert "[Patients]" in out["text"]
    assert "John" in out["text"]


def test_loads_docx_paragraphs_and_tables():
    out = load_document("a.docx", _docx_bytes())
    assert "ICU transfer policy" in out["text"]
    assert "Name | Value" in out["text"]


def test_loads_pdf_with_text_layer():
    out = load_document("a.pdf", _text_pdf("Hello PDF World"))
    assert "Hello PDF World" in out["text"]


# ── rejections ────────────────────────────────────────────────────────────────

def test_rejects_scanned_pdf():
    with pytest.raises(DocumentError) as e:
        load_document("scan.pdf", _blank_pdf())
    assert e.value.code == "scanned-pdf"


def test_rejects_password_protected_pdf():
    with pytest.raises(DocumentError) as e:
        load_document("locked.pdf", _blank_pdf(password="secret"))
    assert e.value.code == "password-protected"


def test_rejects_empty_file():
    with pytest.raises(DocumentError) as e:
        load_document("a.txt", b"")
    assert e.value.code == "empty-file"


def test_rejects_unsupported_extension():
    with pytest.raises(DocumentError) as e:
        load_document("malware.exe", b"MZ\x90\x00 some bytes")
    assert e.value.code == "unsupported-file-type"


def test_rejects_oversized_file(monkeypatch):
    _set_upload_limit_mb(monkeypatch, "0")  # nothing fits
    with pytest.raises(DocumentError) as e:
        load_document("a.txt", b"any content at all")
    assert e.value.code == "file-too-large"


def test_rejects_mime_mismatch():
    # Declared .pdf but the bytes are plain text.
    with pytest.raises(DocumentError) as e:
        load_document("fake.pdf", b"this is not really a pdf, just text")
    assert e.value.code == "unsupported-file-type"


# ── dedup (C-C) ───────────────────────────────────────────────────────────────

def test_duplicate_content_under_different_name_rejected():
    content = b"Identical content body for the dedup check."
    existing = {content_hash(content): "original.txt"}
    with pytest.raises(DocumentError) as e:
        load_document("copy.txt", content, existing_hashes=existing)
    assert e.value.code == "duplicate-content"


def test_same_filename_same_content_is_allowed_replace():
    content = b"Some report content that will be re-uploaded."
    existing = {content_hash(content): "report.txt"}
    out = load_document("report.txt", content, existing_hashes=existing)  # no raise
    assert out["filename"] == "report.txt"


def test_content_hash_is_sha256_first_32_hex():
    content = b"hash me"
    assert content_hash(content) == hashlib.sha256(content).hexdigest()[:32]
    assert len(content_hash(content)) == 32


# ── storage ───────────────────────────────────────────────────────────────────

def test_save_upload_writes_bytes_and_returns_path():
    path = storage.save_upload("note.txt", b"hello bytes")
    with open(path, "rb") as f:
        assert f.read() == b"hello bytes"
